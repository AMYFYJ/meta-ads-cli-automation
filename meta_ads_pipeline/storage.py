from __future__ import annotations

import csv
import os
import shutil
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import Dataset
from .schema import TABLES


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
ID_FILL = PatternFill("solid", fgColor="D9EAD3")
RESULT_FILL = PatternFill("solid", fgColor="FCE4D6")


def load_source(path: str) -> Dataset:
    source = Path(path)
    if source.suffix.lower() == ".xlsx":
        return Dataset(path=str(source), kind="xlsx", tables=_load_xlsx(source))
    if source.suffix.lower() in {".sqlite", ".sqlite3", ".db"}:
        return Dataset(path=str(source), kind="sqlite", tables=_load_sqlite(source))
    if source.is_dir():
        return Dataset(path=str(source), kind="csv_dir", tables=_load_csv_dir(source))
    raise ValueError(f"Unsupported source format: {path}")


def _load_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(path)
    tables: dict[str, list[dict[str, Any]]] = {}
    for sheet_name, spec in TABLES.items():
        if sheet_name not in wb.sheetnames:
            tables[sheet_name] = []
            continue
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        rows: list[dict[str, Any]] = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = {str(header): value for header, value in zip(headers, values) if header}
            if any(value not in (None, "") for value in row.values()):
                rows.append({column: row.get(column, "") for column in spec.columns})
        tables[sheet_name] = rows
    return tables


def _load_sqlite(path: Path) -> dict[str, list[dict[str, Any]]]:
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    try:
        tables: dict[str, list[dict[str, Any]]] = {}
        for table_name, spec in TABLES.items():
            if not _sqlite_table_exists(con, table_name):
                tables[table_name] = []
                continue
            rows = [dict(row) for row in con.execute(f'SELECT * FROM "{table_name}"')]
            tables[table_name] = [{column: row.get(column, "") for column in spec.columns} for row in rows]
        return tables
    finally:
        con.close()


def _sqlite_table_exists(con: sqlite3.Connection, table_name: str) -> bool:
    row = con.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
        (table_name,),
    ).fetchone()
    return row is not None


def _load_csv_dir(path: Path) -> dict[str, list[dict[str, Any]]]:
    tables: dict[str, list[dict[str, Any]]] = {}
    for table_name, spec in TABLES.items():
        csv_path = path / f"{table_name}.csv"
        if not csv_path.exists():
            tables[table_name] = []
            continue
        with csv_path.open(newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)
            tables[table_name] = [{column: row.get(column, "") for column in spec.columns} for row in reader]
    return tables


def create_template(path: str, with_sample: bool = False) -> None:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    sample = sample_tables(output.parent) if with_sample else {}

    for table_name, spec in TABLES.items():
        ws = wb.create_sheet(table_name)
        ws.append(spec.columns)
        for row in sample.get(table_name, []):
            ws.append([row.get(column, "") for column in spec.columns])
        _style_sheet(ws, table_name)
    _add_lookup_sheet(wb)
    wb.save(output)


def create_sqlite_template(path: str, with_sample: bool = False) -> None:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()
    sample = sample_tables(output.parent) if with_sample else {}
    con = sqlite3.connect(output)
    try:
        for table_name, spec in TABLES.items():
            columns = ", ".join(f'"{column}" TEXT' for column in spec.columns)
            con.execute(f'CREATE TABLE "{table_name}" ({columns})')
            for row in sample.get(table_name, []):
                placeholders = ", ".join("?" for _ in spec.columns)
                quoted_cols = ", ".join(f'"{column}"' for column in spec.columns)
                con.execute(
                    f'INSERT INTO "{table_name}" ({quoted_cols}) VALUES ({placeholders})',
                    [row.get(column, "") for column in spec.columns],
                )
        con.commit()
    finally:
        con.close()


def rebuild_workbook(dataset: Dataset, path: str) -> None:
    """Rewrite a dataset as a fresh, fully-styled xlsx whose sheets carry every
    column in the current TABLES schema. Used to migrate older workbooks when
    new columns are added; unknown legacy columns are dropped."""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for table_name, spec in TABLES.items():
        ws = wb.create_sheet(table_name)
        ws.append(spec.columns)
        for row in dataset.tables.get(table_name, []):
            ws.append(["" if row.get(column) is None else row.get(column) for column in spec.columns])
        _style_sheet(ws, table_name)
    _add_lookup_sheet(wb)
    wb.save(output)


def _style_sheet(ws, table_name: str) -> None:
    ws.freeze_panes = "A2"
    max_col = ws.max_column
    max_row = max(ws.max_row, 2)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        header = str(column_cells[0].value)
        width = min(max(len(header) + 4, 14), 34)
        if "copy" in header or "description" in header or "reason" in header or "error" in header:
            width = 38
        ws.column_dimensions[letter].width = width
        if header.startswith("meta_") or header.endswith("_id"):
            for cell in column_cells[1:]:
                cell.fill = ID_FILL
        elif header in {"last_result", "last_error", "result", "error", "updated_at", "applied_at"}:
            for cell in column_cells[1:]:
                cell.fill = RESULT_FILL
        else:
            for cell in column_cells[1:]:
                cell.fill = INPUT_FILL
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    ref = f"A1:{ws.cell(row=max_row, column=max_col).coordinate}"
    tab = Table(displayName=f"{table_name}Table", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)


def _add_lookup_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lookups")
    lookups = [
        ("objective", "OUTCOME_AWARENESS"),
        ("objective", "OUTCOME_TRAFFIC"),
        ("objective", "OUTCOME_ENGAGEMENT"),
        ("objective", "OUTCOME_LEADS"),
        ("objective", "OUTCOME_APP_PROMOTION"),
        ("objective", "OUTCOME_SALES"),
        ("status", "PAUSED"),
        ("status", "ACTIVE"),
        ("budget_mode", "CBO"),
        ("budget_mode", "ABO"),
        ("approval_status", "DRAFT"),
        ("approval_status", "APPROVED"),
        ("approval_status", "REJECTED"),
        ("audience_type", "CUSTOM"),
        ("audience_type", "LOOKALIKE"),
        ("audience_type", "WEBSITE"),
        ("audience_type", "SAVED"),
        ("upload_operation", "ADD"),
        ("upload_operation", "REMOVE"),
        ("upload_operation", "REPLACE"),
        ("automation_toggle", "TRUE"),
        ("automation_toggle", "FALSE"),
        ("format", "image"),
        ("format", "video"),
        ("format", "dco"),
        ("cta", "SHOP_NOW"),
        ("cta", "LEARN_MORE"),
        ("cta", "SIGN_UP"),
        ("cta", "CONTACT_US"),
        ("genders", "all"),
        ("genders", "male"),
        ("genders", "female"),
        ("placements", "automatic"),
        ("placements", "facebook"),
        ("placements", "instagram"),
        ("placements", "audience_network"),
        ("placements", "messenger"),
    ]
    ws.append(["type", "value"])
    for row in lookups:
        ws.append(list(row))
    _style_sheet(ws, "Lookups")


def save_with_updates(
    dataset: Dataset,
    output_path: str,
    updates: list[tuple[str, str, str, Any]],
    append_rows: dict[str, list[dict[str, Any]]] | None = None,
) -> None:
    append_rows = append_rows or {}
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if dataset.kind == "xlsx":
        _save_xlsx_with_updates(Path(dataset.path), output, updates, append_rows)
    elif dataset.kind == "sqlite":
        _save_sqlite_with_updates(Path(dataset.path), output, updates, append_rows)
    elif dataset.kind == "csv_dir":
        _save_csv_dir_with_updates(Path(dataset.path), output, updates, append_rows)
    else:
        raise ValueError(f"Unsupported dataset kind: {dataset.kind}")


def _save_xlsx_with_updates(
    source: Path,
    output: Path,
    updates: list[tuple[str, str, str, Any]],
    append_rows: dict[str, list[dict[str, Any]]],
) -> None:
    shutil.copyfile(source, output)
    wb = load_workbook(output)
    for table_name, rows in append_rows.items():
        if table_name not in wb.sheetnames:
            continue
        ws = wb[table_name]
        headers = [cell.value for cell in ws[1]]
        for row in rows:
            ws.append([row.get(column, "") for column in headers])
    for table_name, key, column, value in updates:
        if table_name not in wb.sheetnames:
            continue
        ws = wb[table_name]
        headers = [cell.value for cell in ws[1]]
        if column not in headers:
            continue
        spec = TABLES[table_name]
        key_column = spec.key_column
        if key_column not in headers:
            continue
        key_idx = headers.index(key_column) + 1
        col_idx = headers.index(column) + 1
        for row_idx in range(2, ws.max_row + 1):
            if str(ws.cell(row=row_idx, column=key_idx).value or "") == str(key):
                ws.cell(row=row_idx, column=col_idx).value = value
                break
    for ws in wb.worksheets:
        _resize_existing_tables(ws)
    wb.save(output)


def _save_sqlite_with_updates(
    source: Path,
    output: Path,
    updates: list[tuple[str, str, str, Any]],
    append_rows: dict[str, list[dict[str, Any]]],
) -> None:
    shutil.copyfile(source, output)
    con = sqlite3.connect(output)
    try:
        for table_name, rows in append_rows.items():
            spec = TABLES[table_name]
            for row in rows:
                columns = [column for column in spec.columns if column in row]
                quoted = ", ".join(f'"{column}"' for column in columns)
                placeholders = ", ".join("?" for _ in columns)
                con.execute(
                    f'INSERT INTO "{table_name}" ({quoted}) VALUES ({placeholders})',
                    [row.get(column, "") for column in columns],
                )
        for table_name, key, column, value in updates:
            spec = TABLES[table_name]
            con.execute(
                f'UPDATE "{table_name}" SET "{column}" = ? WHERE "{spec.key_column}" = ?',
                (value, key),
            )
        con.commit()
    finally:
        con.close()


def _save_csv_dir_with_updates(
    source: Path,
    output: Path,
    updates: list[tuple[str, str, str, Any]],
    append_rows: dict[str, list[dict[str, Any]]],
) -> None:
    if output.exists():
        shutil.rmtree(output)
    shutil.copytree(source, output)
    dataset = load_source(str(output))
    for table_name, key, column, value in updates:
        spec = TABLES[table_name]
        for row in dataset.tables.get(table_name, []):
            if str(row.get(spec.key_column, "")) == str(key):
                row[column] = value
                break
    for table_name, rows in append_rows.items():
        dataset.tables.setdefault(table_name, []).extend(rows)
    for table_name, spec in TABLES.items():
        csv_path = output / f"{table_name}.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=spec.columns)
            writer.writeheader()
            writer.writerows(dataset.tables.get(table_name, []))


def _resize_existing_tables(ws) -> None:
    if not ws.tables:
        return
    max_col = ws.max_column
    max_row = max(ws.max_row, 2)
    ref = f"A1:{get_column_letter(max_col)}{max_row}"
    for table in ws.tables.values():
        table.ref = ref


def sample_tables(base_dir: Path) -> dict[str, list[dict[str, Any]]]:
    asset_dir = base_dir / "assets"
    image_path = asset_dir / "spring_launch_hero.png"
    alt_image_path = asset_dir / "retargeting_offer.png"
    upload_dir = base_dir / "audience_uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    vip_upload_path = upload_dir / "vip_buyers.csv"
    if not vip_upload_path.exists():
        vip_upload_path.write_text(
            "EMAIL,FN,LN\n"
            "hashed_email_1,hashed_first_1,hashed_last_1\n"
            "hashed_email_2,hashed_first_2,hashed_last_2\n",
            encoding="utf-8",
        )
    return {
        "Accounts": [
            {
                "account_key": "acct_demo_us",
                "business_id": "1234567890",
                "ad_account_id": "",
                "create_ad_account": "TRUE",
                "account_name": "Demo DTC Sandbox",
                "currency": "USD",
                "timezone_id": "1",
                "page_id": "112233445566",
                "instagram_actor_id": "223344556677",
                "pixel_dataset_id": "998877665544",
                "catalog_id": "",
                "status": "PLANNED",
                "owner": "Growth Team",
                "max_daily_budget_cents": "25000",
                "approval_required": "TRUE",
            }
        ],
        "Campaigns": [
            {
                "campaign_key": "cmp_spring_launch",
                "account_key": "acct_demo_us",
                "meta_campaign_id": "",
                "name": "Spring Launch | Sales | US",
                "objective": "OUTCOME_SALES",
                "buying_type": "AUCTION",
                "budget_mode": "CBO",
                "daily_budget_cents": "5000",
                "lifetime_budget_cents": "",
                "bid_strategy": "LOWEST_COST_WITHOUT_CAP",
                "special_ad_categories": "[]",
                "desired_status": "PAUSED",
                "launch_batch": "demo_batch_001",
                "approval_status": "APPROVED",
            }
        ],
        "AdSets": [
            {
                "adset_key": "adset_prospecting_us",
                "campaign_key": "cmp_spring_launch",
                "meta_adset_id": "",
                "name": "Prospecting | Broad | US",
                "optimization_goal": "OFFSITE_CONVERSIONS",
                "billing_event": "IMPRESSIONS",
                "daily_budget_cents": "",
                "lifetime_budget_cents": "",
                "bid_amount_cents": "1200",
                "bid_strategy": "",
                "start_time": "2026-06-03T09:00:00-04:00",
                "end_time": "2026-06-17T23:59:00-04:00",
                "countries": "US",
                "regions": "",
                "cities": "",
                "zips": "",
                "age_min": "25",
                "age_max": "54",
                "genders": "all",
                "languages": "",
                "interests": "6003306084421:Yoga, 6003384248805:Fitness and wellness",
                "behaviors": "",
                "flexible_spec_json": "",
                "exclusions_json": "",
                "custom_audiences": "",
                "excluded_audiences": "aud_vip_buyers",
                "placements": "automatic",
                "facebook_positions": "",
                "instagram_positions": "",
                "device_platforms": "",
                "advantage_audience": "TRUE",
                "detailed_targeting_expansion": "TRUE",
                "custom_audience_expansion": "TRUE",
                "advantage_placements": "TRUE",
                "targeting_json": "",
                "targeting_automation_json": "",
                "promoted_object_json": "",
                "pixel_dataset_id": "998877665544",
                "pixel_event": "PURCHASE",
                "attribution_window": "7d_click,1d_view",
                "desired_status": "PAUSED",
                "approval_status": "APPROVED",
            },
            {
                "adset_key": "adset_retargeting_us",
                "campaign_key": "cmp_spring_launch",
                "meta_adset_id": "",
                "name": "Retargeting | Site Visitors | US",
                "optimization_goal": "OFFSITE_CONVERSIONS",
                "billing_event": "IMPRESSIONS",
                "daily_budget_cents": "",
                "lifetime_budget_cents": "",
                "bid_amount_cents": "900",
                "bid_strategy": "",
                "start_time": "2026-06-03T09:00:00-04:00",
                "end_time": "2026-06-17T23:59:00-04:00",
                "countries": "US",
                "regions": "",
                "cities": "",
                "zips": "",
                "age_min": "25",
                "age_max": "64",
                "genders": "all",
                "languages": "",
                "interests": "",
                "behaviors": "",
                "flexible_spec_json": "",
                "exclusions_json": "",
                "custom_audiences": "aud_site_visitors_30d",
                "excluded_audiences": "aud_vip_buyers",
                "placements": "facebook, instagram",
                "facebook_positions": "feed, marketplace",
                "instagram_positions": "stream, story, reels",
                "device_platforms": "mobile, desktop",
                "advantage_audience": "FALSE",
                "detailed_targeting_expansion": "",
                "custom_audience_expansion": "",
                "advantage_placements": "",
                "targeting_json": "",
                "targeting_automation_json": "",
                "promoted_object_json": "",
                "pixel_dataset_id": "998877665544",
                "pixel_event": "PURCHASE",
                "attribution_window": "7d_click,1d_view",
                "desired_status": "PAUSED",
                "approval_status": "APPROVED",
            },
        ],
        "Ads": [
            {
                "ad_key": "ad_prospecting_static_a",
                "adset_key": "adset_prospecting_us",
                "meta_ad_id": "",
                "meta_creative_id": "",
                "name": "Prospecting | Static A",
                "format": "image",
                "asset_path_or_url": os.path.relpath(image_path, base_dir),
                "image_hash_or_video_id": "",
                "primary_text": "Meet the new spring collection built for everyday movement.",
                "headline": "Fresh styles just landed",
                "description": "Limited launch pricing this week.",
                "cta": "SHOP_NOW",
                "destination_url": "https://example.com/spring-launch",
                "url_tags": "utm_source=meta&utm_medium=paid_social",
                "utm_template": "utm_source=meta&utm_medium=paid_social&utm_campaign={campaign_key}&utm_content={ad_key}",
                "page_id": "112233445566",
                "instagram_actor_id": "223344556677",
                "tracking_specs": "",
                "variant_label": "A",
                "desired_status": "PAUSED",
                "launch_batch": "demo_batch_001",
                "approval_status": "APPROVED",
            },
            {
                "ad_key": "ad_retargeting_offer_a",
                "adset_key": "adset_retargeting_us",
                "meta_ad_id": "",
                "meta_creative_id": "",
                "name": "Retargeting | Offer A",
                "format": "image",
                "asset_path_or_url": os.path.relpath(alt_image_path, base_dir),
                "image_hash_or_video_id": "",
                "primary_text": "Still thinking it over? Your launch offer is waiting.",
                "headline": "Come back for 15% off",
                "description": "Complete your order before the launch window closes.",
                "cta": "SHOP_NOW",
                "destination_url": "https://example.com/spring-launch?offer=return",
                "url_tags": "utm_source=meta&utm_medium=paid_social",
                "utm_template": "utm_source=meta&utm_medium=paid_social&utm_campaign={campaign_key}&utm_content={ad_key}",
                "page_id": "112233445566",
                "instagram_actor_id": "223344556677",
                "tracking_specs": "",
                "variant_label": "A",
                "desired_status": "PAUSED",
                "launch_batch": "demo_batch_001",
                "approval_status": "APPROVED",
            },
        ],
        "Audiences": [
            {
                "audience_key": "aud_vip_buyers",
                "account_key": "acct_demo_us",
                "meta_audience_id": "",
                "name": "VIP Buyers Seed",
                "audience_type": "CUSTOM",
                "subtype": "CUSTOM",
                "description": "Seed audience for customers uploaded from CRM.",
                "source_audience_key": "",
                "pixel_dataset_id": "",
                "retention_days": "",
                "countries": "",
                "lookalike_ratio": "",
                "targeting_json": "",
                "rule_json": "",
                "lookalike_spec_json": "",
                "customer_file_source": "USER_PROVIDED_ONLY",
                "upload_operation": "ADD",
                "upload_schema": "EMAIL,FN,LN",
                "upload_data_path": os.path.relpath(vip_upload_path, base_dir),
                "upload_data_json": "",
                "upload_hash_type": "HASHED",
                "approval_status": "APPROVED",
            },
            {
                "audience_key": "aud_vip_lookalike_us",
                "account_key": "acct_demo_us",
                "meta_audience_id": "",
                "name": "VIP Buyers Lookalike 1% US",
                "audience_type": "LOOKALIKE",
                "subtype": "LOOKALIKE",
                "description": "US lookalike based on VIP buyers seed.",
                "source_audience_key": "aud_vip_buyers",
                "pixel_dataset_id": "",
                "retention_days": "",
                "countries": "US",
                "lookalike_ratio": "0.01",
                "targeting_json": "",
                "rule_json": "",
                "lookalike_spec_json": "{\"type\":\"similarity\",\"ratio\":0.01,\"country\":\"US\"}",
                "customer_file_source": "",
                "approval_status": "APPROVED",
            },
            {
                "audience_key": "aud_site_visitors_30d",
                "account_key": "acct_demo_us",
                "meta_audience_id": "",
                "name": "Site Visitors 30D",
                "audience_type": "WEBSITE",
                "subtype": "WEBSITE",
                "description": "Recent website visitors from the demo pixel.",
                "source_audience_key": "",
                "pixel_dataset_id": "998877665544",
                "retention_days": "30",
                "countries": "",
                "lookalike_ratio": "",
                "targeting_json": "",
                "rule_json": "{\"inclusions\":{\"operator\":\"or\",\"rules\":[{\"event_sources\":[{\"id\":\"998877665544\",\"type\":\"pixel\"}],\"retention_seconds\":2592000,\"filter\":{\"operator\":\"and\",\"filters\":[{\"field\":\"event\",\"operator\":\"eq\",\"value\":\"PageView\"}]}}]}}",
                "lookalike_spec_json": "",
                "customer_file_source": "",
                "approval_status": "APPROVED",
            },
            {
                "audience_key": "aud_saved_broad_us",
                "account_key": "acct_demo_us",
                "meta_audience_id": "",
                "name": "Saved Broad US 25-54",
                "audience_type": "SAVED",
                "subtype": "",
                "description": "Reusable broad saved audience.",
                "source_audience_key": "",
                "pixel_dataset_id": "",
                "retention_days": "",
                "countries": "US",
                "lookalike_ratio": "",
                "targeting_json": "{\"geo_locations\":{\"countries\":[\"US\"]},\"age_min\":25,\"age_max\":54}",
                "rule_json": "",
                "lookalike_spec_json": "",
                "customer_file_source": "",
                "approval_status": "APPROVED",
            },
        ],
        "BulkChanges": [
            {
                "change_id": "chg_raise_campaign_budget",
                "operation": "SET_FIELD",
                "object_level": "campaign",
                "object_key_or_meta_id": "cmp_spring_launch",
                "field": "daily_budget_cents",
                "old_value": "5000",
                "new_value": "7500",
                "value_json": "",
                "effective_at": "2026-06-04T09:00:00-04:00",
                "reason": "Increase launch-day budget after creative QA.",
                "requested_by": "Growth Team",
                "approval_status": "APPROVED",
            },
            {
                "change_id": "chg_raise_prospecting_bid_cap",
                "operation": "SET_FIELD",
                "object_level": "adset",
                "object_key_or_meta_id": "adset_prospecting_us",
                "field": "bid_amount_cents",
                "old_value": "1200",
                "new_value": "1500",
                "value_json": "",
                "effective_at": "2026-06-04T09:00:00-04:00",
                "reason": "Raise bid cap to improve launch delivery.",
                "requested_by": "Growth Team",
                "approval_status": "APPROVED",
            },
            {
                "change_id": "chg_extend_retargeting_flight",
                "operation": "SET_FIELD",
                "object_level": "adset",
                "object_key_or_meta_id": "adset_retargeting_us",
                "field": "end_time",
                "old_value": "2026-06-17T23:59:00-04:00",
                "new_value": "2026-06-21T23:59:00-04:00",
                "value_json": "",
                "effective_at": "2026-06-10T09:00:00-04:00",
                "reason": "Extend retargeting window through final promo weekend.",
                "requested_by": "Growth Team",
                "approval_status": "APPROVED",
            },
            {
                "change_id": "chg_patch_retargeting_targeting",
                "operation": "PATCH_JSON",
                "object_level": "adset",
                "object_key_or_meta_id": "adset_retargeting_us",
                "field": "",
                "old_value": "",
                "new_value": "",
                "value_json": "{\"targeting_automation\":{\"advantage_audience\":0}}",
                "effective_at": "2026-06-04T09:00:00-04:00",
                "reason": "Force manual retargeting controls for the demo ad set.",
                "requested_by": "Growth Team",
                "approval_status": "APPROVED",
            },
        ],
        "OptimizationRules": [
            {
                "rule_key": "pause_high_cpa_ads",
                "scope_level": "ad",
                "metric": "cost_per_result",
                "operator": ">",
                "threshold": "40",
                "action_object_level": "ad",
                "action_field": "desired_status",
                "action_value": "PAUSED",
                "max_actions": "10",
                "approval_status": "APPROVED",
                "generated_change_prefix": "opt_pause_high_cpa",
            },
            {
                "rule_key": "raise_budget_high_roas_campaigns",
                "scope_level": "ad",
                "metric": "roas",
                "operator": ">=",
                "threshold": "3",
                "action_object_level": "campaign",
                "action_field": "daily_budget_cents",
                "action_value": "9000",
                "max_actions": "5",
                "approval_status": "APPROVED",
                "generated_change_prefix": "opt_raise_budget",
            },
        ],
    }
