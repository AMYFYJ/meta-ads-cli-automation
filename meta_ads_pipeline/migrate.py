"""Migrate a legacy workbook (separate TargetingPresets / AutomationSettings /
Creatives / AudienceUploads / DuplicateJobs sheets) to the consolidated layout:

- TargetingPresets + adset-level AutomationSettings fold into AdSets columns
- Creatives fold into their Ads rows (creative fields live inline per ad)
- AudienceUploads fold into upload_* columns on Audiences
- Unapplied DuplicateJobs become BulkChanges DUPLICATE rows

Also works on an already-consolidated workbook (it just rewrites the sheets with
the current schema's columns, like the old column-add migration did).
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .planner import _deep_merge  # JSON merge shared with the planner
from .schema import clean
from .storage import load_source, rebuild_workbook

# Preset column -> AdSets column (only filled where the ad set cell is blank).
_PRESET_SIMPLE_COLUMNS = {
    "countries": "countries",
    "regions": "regions",
    "cities": "cities",
    "zips": "zips",
    "age_min": "age_min",
    "age_max": "age_max",
    "genders": "genders",
    "languages": "languages",
    "interests": "interests",
    "behaviors": "behaviors",
    "flexible_spec_json": "flexible_spec_json",
    "exclusions_json": "exclusions_json",
    "facebook_positions": "facebook_positions",
    "instagram_positions": "instagram_positions",
    "device_platforms": "device_platforms",
}

_PRESET_AUDIENCE_COLUMNS = {
    "custom_audience_keys": "custom_audiences",
    "excluded_audience_keys": "excluded_audiences",
}

_AUTOMATION_TOGGLES = (
    "advantage_audience",
    "detailed_targeting_expansion",
    "custom_audience_expansion",
    "advantage_placements",
)

_CREATIVE_COLUMNS = (
    "meta_creative_id",
    "format",
    "asset_path_or_url",
    "image_hash_or_video_id",
    "primary_text",
    "headline",
    "description",
    "cta",
    "destination_url",
    "utm_template",
    "page_id",
    "instagram_actor_id",
    "variant_label",
)

_UPLOAD_COLUMNS = {
    "operation": "upload_operation",
    "schema": "upload_schema",
    "data_path": "upload_data_path",
    "data_json": "upload_data_json",
    "hash_type": "upload_hash_type",
    "applied_at": "upload_applied_at",
    "result": "upload_result",
    "error": "upload_error",
}


def migrate_workbook(source: str, out: str) -> list[str]:
    """Migrate `source` (xlsx) to `out`. Returns human-readable notes."""
    dataset = load_source(source)
    legacy = _load_raw_sheets(
        Path(source),
        ["AdSets", "Ads", "TargetingPresets", "AutomationSettings", "Creatives", "AudienceUploads", "DuplicateJobs"],
    )
    notes: list[str] = []
    _fold_presets(dataset, legacy, notes)
    _fold_automation(dataset, legacy, notes)
    _fold_creatives(dataset, legacy, notes)
    _fold_uploads(dataset, legacy, notes)
    _fold_duplicate_jobs(dataset, legacy, notes)
    rebuild_workbook(dataset, out)
    return notes


def _load_raw_sheets(path: Path, sheet_names: list[str]) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True)
    sheets: dict[str, list[dict[str, Any]]] = {}
    for name in sheet_names:
        if name not in wb.sheetnames:
            sheets[name] = []
            continue
        ws = wb[name]
        rows_iter = ws.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows_iter, [])]
        rows = []
        for values in rows_iter:
            row = {header: value for header, value in zip(headers, values) if header}
            if any(clean(value) for value in row.values()):
                rows.append(row)
        sheets[name] = rows
    wb.close()
    return sheets


def _fold_presets(dataset, legacy, notes: list[str]) -> None:
    presets = {clean(row.get("preset_key")): row for row in legacy["TargetingPresets"]}
    if not presets:
        return
    raw_preset_by_adset = {
        clean(row.get("adset_key")): clean(row.get("targeting_preset_key"))
        for row in legacy["AdSets"]
    }
    used = set()
    for adset in dataset.tables.get("AdSets", []):
        preset_key = raw_preset_by_adset.get(clean(adset.get("adset_key")), "")
        preset = presets.get(preset_key)
        if not preset:
            continue
        used.add(preset_key)
        for source, target in _PRESET_SIMPLE_COLUMNS.items():
            if not clean(adset.get(target)) and clean(preset.get(source)):
                adset[target] = clean(preset.get(source))
        # The consolidated `placements` column is the publisher-platform list.
        preset_placements = clean(preset.get("publisher_platforms")) or clean(preset.get("placements"))
        if not clean(adset.get("placements")) and preset_placements:
            adset["placements"] = preset_placements
        for source, target in _PRESET_AUDIENCE_COLUMNS.items():
            merged = _merge_csv(preset.get(source), adset.get(target))
            if merged:
                adset[target] = merged
        adset["targeting_json"] = _merge_json_text(preset.get("targeting_json"), adset.get("targeting_json"))
        notes.append(f"AdSets[{clean(adset.get('adset_key'))}]: folded targeting preset '{preset_key}' into columns.")
    for preset_key in sorted(set(presets) - used):
        notes.append(f"TargetingPresets[{preset_key}]: unused preset dropped (presets no longer have a tab).")


def _fold_automation(dataset, legacy, notes: list[str]) -> None:
    adsets = {clean(row.get("adset_key")): row for row in dataset.tables.get("AdSets", [])}
    for row in legacy["AutomationSettings"]:
        key = clean(row.get("setting_key"))
        level = clean(row.get("object_level")).lower()
        target = adsets.get(clean(row.get("object_key"))) if level == "adset" else None
        if target is None:
            notes.append(f"AutomationSettings[{key}]: dropped ({level}-level automation has no consolidated home).")
            continue
        for toggle in _AUTOMATION_TOGGLES:
            if not clean(target.get(toggle)) and clean(row.get(toggle)):
                target[toggle] = clean(row.get(toggle))
        target["targeting_automation_json"] = _merge_json_text(
            row.get("targeting_automation_json"), target.get("targeting_automation_json")
        )
        notes.append(f"AdSets[{clean(target.get('adset_key'))}]: folded automation settings '{key}' into toggle columns.")


def _fold_creatives(dataset, legacy, notes: list[str]) -> None:
    creatives = {clean(row.get("creative_key")): row for row in legacy["Creatives"]}
    if not creatives:
        return
    raw_creative_by_ad = {
        clean(row.get("ad_key")): clean(row.get("creative_key"))
        for row in legacy["Ads"]
    }
    referenced: dict[str, list[str]] = {}
    for ad in dataset.tables.get("Ads", []):
        ad_key = clean(ad.get("ad_key"))
        creative_key = raw_creative_by_ad.get(ad_key, "")
        creative = creatives.get(creative_key)
        if not creative:
            if creative_key:
                notes.append(f"Ads[{ad_key}]: creative_key '{creative_key}' had no Creatives row; creative fields left blank.")
            continue
        referenced.setdefault(creative_key, []).append(ad_key)
        for column in _CREATIVE_COLUMNS:
            if not clean(ad.get(column)) and clean(creative.get(column)):
                ad[column] = clean(creative.get(column))
        for dropped in ("compliance_notes", "creative_angle"):
            if clean(creative.get(dropped)):
                notes.append(f"Ads[{ad_key}]: creative {dropped} not carried over: {clean(creative.get(dropped))}")
    for creative_key, ad_keys in referenced.items():
        if len(ad_keys) > 1 and not clean(creatives[creative_key].get("meta_creative_id")):
            notes.append(
                f"Creative '{creative_key}' was shared by ads {', '.join(ad_keys)} and has no Meta ID yet; "
                "each ad row now builds its own creative unless you paste one meta_creative_id into the others."
            )
    for creative_key in sorted(set(creatives) - set(referenced) - {""}):
        notes.append(f"Creatives[{creative_key}]: not referenced by any ad; dropped.")


def _fold_uploads(dataset, legacy, notes: list[str]) -> None:
    audiences = {clean(row.get("audience_key")): row for row in dataset.tables.get("Audiences", [])}
    seen: set[str] = set()
    for row in legacy["AudienceUploads"]:
        upload_key = clean(row.get("upload_key"))
        audience = audiences.get(clean(row.get("audience_key")))
        if audience is None:
            notes.append(f"AudienceUploads[{upload_key}]: no matching audience; dropped.")
            continue
        audience_key = clean(row.get("audience_key"))
        if audience_key in seen:
            notes.append(f"AudienceUploads[{upload_key}]: audience '{audience_key}' already carries an upload; extra row dropped.")
            continue
        seen.add(audience_key)
        for source, target in _UPLOAD_COLUMNS.items():
            if clean(row.get(source)):
                audience[target] = clean(row.get(source))
        notes.append(f"Audiences[{audience_key}]: folded upload '{upload_key}' into upload_* columns.")


def _fold_duplicate_jobs(dataset, legacy, notes: list[str]) -> None:
    for row in legacy["DuplicateJobs"]:
        key = clean(row.get("job_key"))
        if clean(row.get("applied_at")):
            notes.append(f"DuplicateJobs[{key}]: already applied; history row dropped (see PublishLog).")
            continue
        overrides = {
            "deep_copy": clean(row.get("deep_copy")),
            "status_option": clean(row.get("status_option")),
            "rename_options": {
                "rename_strategy": clean(row.get("rename_strategy")),
                "prefix": clean(row.get("name_prefix")),
                "suffix": clean(row.get("name_suffix")),
            },
        }
        overrides["rename_options"] = {k: v for k, v in overrides["rename_options"].items() if v}
        overrides = {k: v for k, v in overrides.items() if v}
        overrides = _deep_merge(overrides, _json_text(row.get("overrides_json")))
        dataset.tables.setdefault("BulkChanges", []).append(
            {
                "change_id": f"chg_migrated_{key}",
                "operation": "DUPLICATE",
                "object_level": clean(row.get("object_level")),
                "object_key_or_meta_id": clean(row.get("source_key_or_meta_id")),
                "field": "",
                "old_value": "",
                "new_value": clean(row.get("copy_count")) or "1",
                "value_json": json.dumps(overrides) if overrides else "",
                "effective_at": "",
                "reason": f"Migrated from DuplicateJobs row {key}.",
                "requested_by": "migration",
                "approval_status": clean(row.get("approval_status")),
                "applied_at": "",
                "result": "",
                "error": "",
            }
        )
        notes.append(f"DuplicateJobs[{key}]: converted to BulkChanges row chg_migrated_{key} (operation=DUPLICATE).")


def _merge_csv(*values: Any) -> str:
    parts: list[str] = []
    for value in values:
        for part in clean(value).split(","):
            part = part.strip()
            if part and part not in parts:
                parts.append(part)
    return ", ".join(parts)


def _json_text(value: Any) -> dict[str, Any]:
    text = clean(value)
    if not text:
        return {}
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _merge_json_text(base: Any, override: Any) -> str:
    merged = _deep_merge(_json_text(base), _json_text(override))
    return json.dumps(merged) if merged else clean(override) or clean(base)
