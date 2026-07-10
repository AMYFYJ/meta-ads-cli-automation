from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from meta_ads_pipeline.assets import ensure_sample_assets
from meta_ads_pipeline.executor import execute_actions
from meta_ads_pipeline.insights import generate_mock_insights
from meta_ads_pipeline.planner import build_plan
from meta_ads_pipeline.storage import create_sqlite_template, create_template, load_source, save_with_updates
from meta_ads_pipeline.validators import has_blocking_errors, validate_dataset


class PipelineTest(unittest.TestCase):
    def test_sample_workbook_plans_and_applies_end_to_end(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "template.xlsx"
            applied = root / "applied.xlsx"
            state = root / "state.json"
            ensure_sample_assets(str(root))
            create_template(str(workbook), with_sample=True)

            dataset = load_source(str(workbook))
            issues = validate_dataset(dataset)
            self.assertFalse(has_blocking_errors(issues), [issue.message for issue in issues])

            actions = build_plan(dataset)
            self.assertEqual(len(actions), 20)
            results, updates, append_rows = execute_actions(dataset, actions, mode="mock", state_path=str(state))
            self.assertTrue(all(result.ok for result in results), [result.message for result in results])
            save_with_updates(dataset, str(applied), updates, append_rows)

            applied_dataset = load_source(str(applied))
            self.assertEqual(build_plan(applied_dataset), [])
            ids = {
                "account": applied_dataset.tables["Accounts"][0]["ad_account_id"],
                "audience": applied_dataset.tables["Audiences"][0]["meta_audience_id"],
                "campaign": applied_dataset.tables["Campaigns"][0]["meta_campaign_id"],
                "ad": applied_dataset.tables["Ads"][0]["meta_ad_id"],
            }
            self.assertTrue(ids["account"].startswith("act_mock_"))
            self.assertTrue(ids["audience"].startswith("mock_aud_"))
            self.assertTrue(ids["campaign"].startswith("mock_cmp_"))
            self.assertTrue(ids["ad"].startswith("mock_ad_"))
            self.assertEqual(applied_dataset.tables["Campaigns"][0]["daily_budget_cents"], "7500")
            self.assertEqual(applied_dataset.tables["AdSets"][0]["bid_amount_cents"], "1500")
            self.assertEqual(applied_dataset.tables["BulkChanges"][0]["result"], "MOCK_UPDATED")
            self.assertTrue(applied_dataset.tables["AudienceUploads"][0]["result"].startswith("MOCK_UPLOADED_"))
            self.assertEqual(applied_dataset.tables["DuplicateJobs"][0]["result"], "MOCK_DUPLICATED")

            state_data = json.loads(state.read_text(encoding="utf-8"))
            self.assertIn("cmp_spring_launch", state_data["ids"]["campaign"])
            self.assertIn("aud_vip_buyers", state_data["ids"]["audience"])

    def test_mock_insights_append_rows_for_created_ads(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "template.xlsx"
            applied = root / "applied.xlsx"
            insights = root / "insights.xlsx"
            state = root / "state.json"
            ensure_sample_assets(str(root))
            create_template(str(workbook), with_sample=True)
            dataset = load_source(str(workbook))
            actions = build_plan(dataset)
            _, updates, append_rows = execute_actions(dataset, actions, mode="mock", state_path=str(state))
            save_with_updates(dataset, str(applied), updates, append_rows)

            applied_dataset = load_source(str(applied))
            save_with_updates(applied_dataset, str(insights), [], generate_mock_insights(applied_dataset))
            wb = load_workbook(insights, data_only=True)
            self.assertEqual(wb["PerformanceSnapshots"].max_row, 3)

    def test_live_cli_commands_place_ad_account_after_ads_group(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "template.xlsx"
            ensure_sample_assets(str(root))
            create_template(str(workbook), with_sample=True)

            dataset = load_source(str(workbook))
            account = dataset.tables["Accounts"][0]
            account["ad_account_id"] = "act_123"
            account["create_ad_account"] = "FALSE"

            campaign_action = next(
                action for action in build_plan(dataset) if action.action_id == "010_create_campaign_cmp_spring_launch"
            )
            self.assertEqual(
                campaign_action.command[:7],
                ["meta", "--output", "json", "--no-input", "ads", "--ad-account-id", "act_123"],
            )

    def test_campaign_bid_strategy_carried_via_graph_followup(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "template.xlsx"
            ensure_sample_assets(str(root))
            create_template(str(workbook), with_sample=True)

            dataset = load_source(str(workbook))
            campaign = dataset.tables["Campaigns"][0]
            campaign["bid_strategy"] = "LOWEST_COST_WITHOUT_CAP"

            actions = build_plan(dataset)
            create = next(action for action in actions if action.action_id == "010_create_campaign_cmp_spring_launch")
            followup = next(
                action for action in actions if action.action_id == "011_set_campaign_bid_strategy_cmp_spring_launch"
            )
            self.assertEqual(followup.executor, "graph")
            self.assertEqual(followup.method, "POST")
            self.assertEqual(followup.endpoint, "${campaign:cmp_spring_launch}")
            self.assertEqual(followup.body, {"bid_strategy": "LOWEST_COST_WITHOUT_CAP"})
            self.assertEqual(followup.depends_on, ["010_create_campaign_cmp_spring_launch"])
            self.assertGreater(actions.index(followup), actions.index(create))

            campaign["bid_strategy"] = ""
            defaulted = next(
                action
                for action in build_plan(dataset)
                if action.action_id == "011_set_campaign_bid_strategy_cmp_spring_launch"
            )
            self.assertEqual(defaulted.body, {"bid_strategy": "LOWEST_COST_WITHOUT_CAP"})

            campaign["bid_strategy"] = "COST_CAP"
            explicit = next(
                action
                for action in build_plan(dataset)
                if action.action_id == "011_set_campaign_bid_strategy_cmp_spring_launch"
            )
            self.assertEqual(explicit.body, {"bid_strategy": "COST_CAP"})

            campaign["bid_strategy"] = ""
            campaign["budget_mode"] = "ABO"
            abo_actions = build_plan(dataset)
            self.assertFalse(
                [action for action in abo_actions if action.action_id.startswith("011_set_campaign_bid_strategy_")]
            )

    def test_chained_applies_keep_publish_log_ids_unique(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "template.xlsx"
            applied = root / "applied.xlsx"
            reapplied = root / "reapplied.xlsx"
            state = root / "state.json"
            ensure_sample_assets(str(root))
            create_template(str(workbook), with_sample=True)
            dataset = load_source(str(workbook))
            _, updates, append_rows = execute_actions(dataset, build_plan(dataset), mode="mock", state_path=str(state))
            save_with_updates(dataset, str(applied), updates, append_rows)

            applied_dataset = load_source(str(applied))
            applied_dataset.tables["Ads"][0]["meta_ad_id"] = ""
            actions = build_plan(applied_dataset)
            self.assertTrue(actions)
            _, updates, append_rows = execute_actions(applied_dataset, actions, mode="mock", state_path=str(state))
            save_with_updates(applied_dataset, str(reapplied), updates, append_rows)

            log_ids = [row["log_id"] for row in load_source(str(reapplied)).tables["PublishLog"]]
            self.assertEqual(len(log_ids), len(set(log_ids)))
            self.assertFalse(has_blocking_errors(validate_dataset(load_source(str(reapplied)))))

    def test_sqlite_template_uses_same_schema(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            db_path = root / "template.sqlite"
            ensure_sample_assets(str(root))
            create_sqlite_template(str(db_path), with_sample=True)
            dataset = load_source(str(db_path))
            self.assertEqual(len(dataset.tables["Campaigns"]), 1)
            self.assertEqual(len(dataset.tables["Audiences"]), 4)
            self.assertFalse(has_blocking_errors(validate_dataset(dataset)))


if __name__ == "__main__":
    unittest.main()
