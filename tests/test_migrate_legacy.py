"""migrate_workbook folds the retired tabs (TargetingPresets, AutomationSettings,
Creatives, AudienceUploads, DuplicateJobs) into the consolidated sheets."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from meta_ads_pipeline.migrate import migrate_workbook
from meta_ads_pipeline.storage import load_source


def _sheet(wb: Workbook, name: str, rows: list[dict[str, str]]) -> None:
    ws = wb.create_sheet(name)
    headers = list(rows[0])
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def _legacy_workbook(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Accounts", [{
        "account_key": "acct", "ad_account_id": "act_1", "account_name": "A",
        "currency": "USD", "timezone_id": "1",
    }])
    _sheet(wb, "Campaigns", [{
        "campaign_key": "cmp", "account_key": "acct", "name": "C", "objective": "OUTCOME_SALES",
        "budget_mode": "CBO", "daily_budget_cents": "5000", "desired_status": "PAUSED",
        "approval_status": "APPROVED",
    }])
    _sheet(wb, "AdSets", [{
        "adset_key": "as1", "campaign_key": "cmp", "name": "AS", "optimization_goal": "LINK_CLICKS",
        "billing_event": "IMPRESSIONS", "targeting_preset_key": "tp1", "countries": "",
        "age_min": "", "genders": "", "desired_status": "PAUSED", "approval_status": "APPROVED",
    }])
    _sheet(wb, "TargetingPresets", [{
        "preset_key": "tp1", "account_key": "acct", "name": "P", "countries": "US, CA",
        "age_min": "21", "age_max": "45", "genders": "female",
        "interests": "6003306084421:Yoga", "custom_audience_keys": "aud1",
        "publisher_platforms": "facebook,instagram", "facebook_positions": "feed",
        "targeting_json": '{"locales": [6]}', "approval_status": "APPROVED",
    }])
    _sheet(wb, "AutomationSettings", [{
        "setting_key": "auto1", "object_level": "adset", "object_key": "as1",
        "advantage_audience": "TRUE", "advantage_placements": "FALSE",
        "targeting_automation_json": "", "approval_status": "APPROVED",
    }])
    _sheet(wb, "Creatives", [{
        "creative_key": "cr1", "account_key": "acct", "meta_creative_id": "998877",
        "name": "Creative", "format": "image", "asset_path_or_url": "",
        "primary_text": "Body", "headline": "Head", "cta": "SHOP_NOW",
        "destination_url": "https://example.com", "page_id": "42",
        "compliance_notes": "note!", "approval_status": "APPROVED",
    }])
    _sheet(wb, "Ads", [{
        "ad_key": "ad1", "adset_key": "as1", "creative_key": "cr1", "name": "Ad",
        "desired_status": "PAUSED", "approval_status": "APPROVED",
    }])
    _sheet(wb, "Audiences", [{
        "audience_key": "aud1", "account_key": "acct", "name": "Aud",
        "audience_type": "CUSTOM", "customer_file_source": "USER_PROVIDED_ONLY",
        "approval_status": "APPROVED",
    }])
    _sheet(wb, "AudienceUploads", [{
        "upload_key": "up1", "audience_key": "aud1", "operation": "ADD",
        "schema": "EMAIL", "data_json": '[["hashed"]]', "hash_type": "HASHED",
        "approval_status": "APPROVED",
    }])
    _sheet(wb, "DuplicateJobs", [{
        "job_key": "dj1", "object_level": "adset", "source_key_or_meta_id": "as1",
        "copy_count": "2", "deep_copy": "TRUE", "status_option": "PAUSED",
        "approval_status": "APPROVED", "applied_at": "",
    }])
    wb.save(path)


class MigrateLegacyWorkbookTest(unittest.TestCase):
    def test_legacy_tabs_fold_into_consolidated_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            legacy = Path(tmp) / "legacy.xlsx"
            migrated = Path(tmp) / "migrated.xlsx"
            _legacy_workbook(legacy)

            notes = migrate_workbook(str(legacy), str(migrated))
            ds = load_source(str(migrated))

            adset = ds.tables["AdSets"][0]
            self.assertEqual(adset["countries"], "US, CA")
            self.assertEqual(adset["age_min"], "21")
            self.assertEqual(adset["genders"], "female")
            self.assertEqual(adset["interests"], "6003306084421:Yoga")
            self.assertEqual(adset["custom_audiences"], "aud1")
            self.assertEqual(adset["placements"], "facebook,instagram")
            self.assertEqual(adset["facebook_positions"], "feed")
            self.assertIn("locales", str(adset["targeting_json"]))
            self.assertEqual(adset["advantage_audience"], "TRUE")

            ad = ds.tables["Ads"][0]
            self.assertEqual(ad["meta_creative_id"], "998877")
            self.assertEqual(ad["primary_text"], "Body")
            self.assertEqual(ad["headline"], "Head")
            self.assertEqual(ad["destination_url"], "https://example.com")

            audience = ds.tables["Audiences"][0]
            self.assertEqual(audience["upload_operation"], "ADD")
            self.assertEqual(audience["upload_schema"], "EMAIL")
            self.assertEqual(audience["upload_data_json"], '[["hashed"]]')

            dup = next(r for r in ds.tables["BulkChanges"] if r["change_id"] == "chg_migrated_dj1")
            self.assertEqual(dup["operation"], "DUPLICATE")
            self.assertEqual(dup["object_key_or_meta_id"], "as1")
            self.assertEqual(dup["new_value"], "2")

            self.assertTrue(any("compliance_notes" in note for note in notes))


if __name__ == "__main__":
    unittest.main()
