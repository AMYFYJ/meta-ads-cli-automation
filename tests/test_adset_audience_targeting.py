"""Ad set audience targeting: workbook columns -> Graph targeting payloads."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from meta_ads_pipeline.assets import ensure_sample_assets
from meta_ads_pipeline.planner import build_plan
from meta_ads_pipeline.schema import TABLES
from meta_ads_pipeline.storage import create_template, load_source, rebuild_workbook
from meta_ads_pipeline.validators import validate_dataset


def _sample_dataset(root: Path):
    workbook = root / "template.xlsx"
    ensure_sample_assets(str(root))
    create_template(str(workbook), with_sample=True)
    return load_source(str(workbook))


def _adset_action(actions, key):
    return next(
        action for action in actions
        if action.object_type == "adset" and action.operation == "create" and action.object_key == key
    )


class AdSetAudienceTargetingTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dataset = _sample_dataset(Path(self._tmp.name))
        self.adset = self.dataset.tables["AdSets"][0]

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def targeting(self):
        actions = build_plan(self.dataset)
        return _adset_action(actions, self.adset["adset_key"]).body["targeting"]

    def test_age_genders_go_through_graph_even_without_preset_or_json(self) -> None:
        self.adset.update({"interests": "", "genders": "female", "age_min": "21", "age_max": "34"})
        actions = build_plan(self.dataset)
        action = _adset_action(actions, self.adset["adset_key"])
        self.assertEqual(action.executor, "graph")
        self.assertEqual(action.body["targeting"]["age_min"], "21")
        self.assertEqual(action.body["targeting"]["age_max"], "34")
        self.assertEqual(action.body["targeting"]["genders"], [2])

    def test_location_columns_build_geo_locations(self) -> None:
        self.adset.update({"countries": "US, CA", "regions": "3847", "cities": "2418779", "zips": "US:94025"})
        targeting = self.targeting()
        self.assertEqual(targeting["geo_locations"]["countries"], ["US", "CA"])
        self.assertEqual(targeting["geo_locations"]["regions"], [{"key": "3847"}])
        self.assertEqual(targeting["geo_locations"]["cities"], [{"key": "2418779"}])
        self.assertEqual(targeting["geo_locations"]["zips"], [{"key": "US:94025"}])

    def test_interests_and_behaviors_parse_id_name_pairs(self) -> None:
        self.adset["interests"] = "6003306084421:Yoga, 6003397496347"
        self.adset["behaviors"] = "6002714895372:Frequent travelers"
        targeting = self.targeting()
        self.assertEqual(
            targeting["interests"],
            [{"id": "6003306084421", "name": "Yoga"}, {"id": "6003397496347"}],
        )
        self.assertEqual(targeting["behaviors"], [{"id": "6002714895372", "name": "Frequent travelers"}])

    def test_flexible_spec_and_exclusions_json(self) -> None:
        self.adset["flexible_spec_json"] = (
            '[{"interests": [{"id": "6003306084421", "name": "Yoga"}]},'
            ' {"behaviors": [{"id": "6002714895372", "name": "Frequent travelers"}]}]'
        )
        self.adset["exclusions_json"] = '{"interests": [{"id": "6003397496347"}]}'
        targeting = self.targeting()
        self.assertEqual(len(targeting["flexible_spec"]), 2)
        self.assertEqual(targeting["exclusions"], {"interests": [{"id": "6003397496347"}]})

    def test_custom_audience_keys_resolve_to_meta_id_or_placeholder(self) -> None:
        self.adset["custom_audiences"] = "aud_site_visitors_30d, 23851234567890123"
        self.adset["excluded_audiences"] = "aud_vip_buyers"
        for audience in self.dataset.tables["Audiences"]:
            if audience["audience_key"] == "aud_vip_buyers":
                audience["meta_audience_id"] = "9988776655"
        actions = build_plan(self.dataset)
        action = _adset_action(actions, self.adset["adset_key"])
        targeting = action.body["targeting"]
        self.assertEqual(
            targeting["custom_audiences"],
            [{"id": "${audience:aud_site_visitors_30d}"}, {"id": "23851234567890123"}],
        )
        self.assertEqual(targeting["excluded_custom_audiences"], [{"id": "9988776655"}])
        self.assertIn("015_create_audience_aud_site_visitors_30d", action.depends_on)
        self.assertNotIn("015_create_audience_aud_vip_buyers", action.depends_on)

    def test_targeting_json_merges_over_columns(self) -> None:
        self.adset["targeting_json"] = '{"age_min": 30, "geo_locations": {"location_types": ["home"]}}'
        targeting = self.targeting()
        self.assertEqual(targeting["age_min"], 30)  # raw JSON wins over the column
        self.assertEqual(targeting["age_max"], "54")  # columns fill the rest
        self.assertEqual(targeting["geo_locations"]["countries"], ["US"])
        self.assertEqual(targeting["geo_locations"]["location_types"], ["home"])


class AdSetTargetingValidationTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dataset = _sample_dataset(Path(self._tmp.name))
        self.adset = self.dataset.tables["AdSets"][0]

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def errors_for(self, field):
        return [
            issue for issue in validate_dataset(self.dataset)
            if issue.table == "AdSets" and issue.field == field and issue.severity == "ERROR"
        ]

    def test_age_bounds(self) -> None:
        self.adset["age_min"] = "12"
        self.assertTrue(self.errors_for("age_min"))
        self.adset["age_min"] = "40"
        self.adset["age_max"] = "30"
        self.assertTrue(self.errors_for("age_max"))

    def test_gender_values(self) -> None:
        self.adset["genders"] = "women"
        self.assertTrue(self.errors_for("genders"))

    def test_interest_entries_need_numeric_ids(self) -> None:
        self.adset["interests"] = "Yoga"
        errors = self.errors_for("interests")
        self.assertTrue(errors)
        self.assertIn("targeting-search", errors[0].message)

    def test_unknown_audience_reference_is_blocking(self) -> None:
        self.adset["custom_audiences"] = "not_a_key_or_id"
        self.assertTrue(self.errors_for("custom_audiences"))

    def test_flexible_spec_must_be_json_list_or_object(self) -> None:
        self.adset["flexible_spec_json"] = '"just a string"'
        self.assertTrue(self.errors_for("flexible_spec_json"))

    def test_valid_sample_has_no_targeting_errors(self) -> None:
        for field in ("age_min", "age_max", "genders", "interests", "behaviors", "custom_audiences", "excluded_audiences"):
            self.assertFalse(self.errors_for(field), field)


class WorkbookMigrationTest(unittest.TestCase):
    def test_rebuild_adds_new_columns_and_keeps_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dataset = _sample_dataset(root)
            migrated_path = root / "migrated.xlsx"
            rebuild_workbook(dataset, str(migrated_path))
            migrated = load_source(str(migrated_path))
            from openpyxl import load_workbook

            headers = [cell.value for cell in load_workbook(migrated_path)["AdSets"][1]]
            self.assertEqual(headers, TABLES["AdSets"].columns)
            for table_name in ("Accounts", "Campaigns", "AdSets", "Ads", "Audiences"):
                self.assertEqual(
                    len(migrated.tables[table_name]), len(dataset.tables[table_name]), table_name
                )


if __name__ == "__main__":
    unittest.main()
